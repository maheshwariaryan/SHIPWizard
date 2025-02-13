from fastapi import FastAPI, HTTPException
from pydantic import BaseModel
from docx import Document

app = FastAPI()

class TextReplacement(BaseModel):
    replacement_text: str

# Function to replace text in the document
def replace_placeholder(doc, placeholder, replacement):
    for para in doc.paragraphs:
        if placeholder in para.text:
            para.text = para.text.replace(placeholder, replacement)

@app.post("/update-doc/")
def update_doc(text_replacement: TextReplacement):
    replacement_text = text_replacement.replacement_text
    
    # Open the existing Word document
    doc = Document('template.docx')
    
    # Replace specific placeholders with desired text
    replace_placeholder(doc, '{text}', replacement_text)
    
    # Save the modified document
    doc.save('C:/Users/mahes/Downloads/Coding/frontend hackathon (real)/ai-interface/public/filled_template.docx')
    
    return {"message": "Document updated successfully"}

if __name__ == '__main__':
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8003)